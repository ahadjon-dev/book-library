import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

DOCS_DIR = Path("docs")
IMAGES_DIR = DOCS_DIR / "images"

def create_docx():
    doc = Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    title = doc.add_paragraph()
    r = title.add_run("📖 Book Library — User Guide & Quick Start")
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor(24, 24, 27)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # Intro
    p = doc.add_paragraph("Welcome to your personal and household Book Library! Whether you have 20 books or 2,000, this app helps you catalog your physical collection, track what you are reading, lend books to friends without losing them, and decide what to read next.")
    p.runs[0].font.size = Pt(11)

    def add_h1(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = RGBColor(14, 116, 144) # Teal / Cyan accent
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)

    def add_h2(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(51, 65, 85)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)

    def add_img(filename, width_in=5.5):
        img_path = IMAGES_DIR / filename
        if img_path.exists():
            doc.add_paragraph()
            doc.add_picture(str(img_path), width=Inches(width_in))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()

    # Section 1
    add_h1("1. Viewing & Exploring Your Library")
    add_img("gallery_view.png", width_in=5.8)
    doc.add_paragraph("You can view and browse your collection in two ways:")
    
    add_h2("🖼️ Gallery View (Visual Grid)")
    doc.add_paragraph("Displays your books as cards with high-resolution cover art. Directly on any book card:")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Click the Status Badge (e.g. ⏳ Unread ▾) to instantly mark it as 📖 Reading or ✅ Finished.")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Click the Stars (★) to rate the book from 1 to 10 without leaving the page.")

    add_h2("📋 Table View (Spreadsheet Style)")
    doc.add_paragraph("Perfect for quickly sorting your library by Author, Publication Year, Page Count, or Shelf Location.")

    add_h2("🏷️ Shelves & Tags (Finding Where Books Are)")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Shelves: Physical locations in your home (e.g., 'Living Room Shelf A', 'Bedside Table', 'Office Desk').")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Tags: Themes, topics, or genres (e.g., 'Favorites', 'Sci-Fi', 'Clean Code', 'Gift').")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Use the Filter Sidebar on the left to instantly filter books by genre, shelf, or author with one click.")

    # Section 2
    add_h1("2. Adding Books (4 Fast Ways)")
    doc.add_paragraph("Click the '+ Add Books to Library' button in the top navigation bar to open the add modal:")

    add_h2("📸 Method A: Take a Photo of Your Bookshelf")
    doc.add_paragraph("Got a full shelf of books? You don't need to type them one by one.")
    add_img("add_shelf_photo.png", width_in=5.0)
    for step in [
        "1. Select the Shelf Photo tab.",
        "2. Snap or upload a photo of your physical bookshelf.",
        "3. The AI reads the book spines, finds cover art, authors, and page counts automatically.",
        "4. Review the detected list and click Add Selected Books."
    ]:
        doc.add_paragraph(step)

    add_h2("📷 Method B: Scan Barcode or Look Up by ISBN")
    doc.add_paragraph("Have a book in your hand?")
    add_img("add_barcode_isbn.png", width_in=5.0)
    for step in [
        "1. Select the Barcode / ISBN tab.",
        "2. Click Scan to point your phone or webcam camera at the barcode on the back cover.",
        "3. Or type the ISBN (the 10 or 13-digit number) and click Look up.",
        "4. The title, author, publisher, cover image, and page count will auto-populate instantly. Click Add Book to Library."
    ]:
        doc.add_paragraph(step)

    add_h2("✍️ Method C: Manual Entry")
    doc.add_paragraph("For rare, vintage, or custom books without barcodes.")
    add_img("add_manual_entry.png", width_in=5.0)
    for step in [
        "1. Select the Manual Entry tab.",
        "2. Fill in the title, author(s), shelf, genre, and page count.",
        "3. Optionally upload a custom cover photo and click Add Book to Library."
    ]:
        doc.add_paragraph(step)

    add_h2("📁 Method D: CSV / Goodreads Import")
    doc.add_paragraph("Transfer your entire reading history from Goodreads or an Excel spreadsheet in one go.")
    add_img("add_csv_goodreads.png", width_in=5.0)
    for step in [
        "1. Select the CSV / Goodreads tab.",
        "2. Drag and drop your .csv file.",
        "3. Click Import CSV — your books, ratings, and read dates will import in seconds."
    ]:
        doc.add_paragraph(step)

    # Section 3
    add_h1("3. Tracking Your Reading Progress & Ratings")
    doc.add_paragraph("Every book has a reading status: ⏳ Unread, 📖 Reading, ✅ Finished, or 🛑 Abandoned.")
    doc.add_paragraph("⚡ 1-Tap Quick Actions: You don't have to open a book detail page to update your progress. Click the status badge or star rating directly on the card.")

    # Section 4
    add_h1("4. Sharing with Family & Roommates (Household Library)")
    doc.add_paragraph("Do you share your physical home library with your spouse, family, or roommate? With Household Libraries, everyone shares the catalog of physical books, but maintains independent reading progress and private ratings.")
    doc.add_paragraph("How to Invite a Member:")
    for step in [
        "1. Click your profile menu in the top right -> Household Library.",
        "2. Copy your 6-character Invite Code or invite link.",
        "3. Your family member creates an account and enters the code to join your library."
    ]:
        doc.add_paragraph(step)

    # Section 5
    add_h1("5. Lending Tracker — Never Lose a Book")
    doc.add_paragraph("Keep track of books you lend to friends or colleagues:")
    for step in [
        "1. Go to the Book Loans tab in the top navigation.",
        "2. Click '+ Lend a Book'.",
        "3. Choose the book, enter your friend's name, contact details, and a return due date.",
        "4. The book will display an orange 'On Loan' badge across your library.",
        "5. When returned, click 'Mark as Returned' with 1 tap."
    ]:
        doc.add_paragraph(step)

    # Section 6
    add_h1("6. Reading Goals & Personal Stats")
    doc.add_paragraph("Stay motivated with your yearly reading challenge. Go to the Stats & Goals tab and set your Yearly Reading Goal (e.g. 25 books in 2026). The app tracks your live pace: Ahead of schedule, On track, or Behind pace.")

    # Section 7
    add_h1("7. AI 'What to Read Next?' Assistant")
    doc.add_paragraph("Standing in front of your shelf wondering what to read? Click '✨ What to Read Next?' in the top navigation, select your mood ('Quick Weekend Read', 'Deep & Thoughtful', 'Fast-Paced'), and the assistant recommends top matches from your unread shelf.")

    # Section 8
    add_h1("8. Public Shareable Profile")
    doc.add_paragraph("Want to show friends what books you own? Click 'Share Your Shelf', toggle Public Shelf to ON, and share your custom link (e.g. /share/your-name). Anyone can browse your shelf without creating an account.")

    # Section 9
    add_h1("9. Exporting Your Data")
    doc.add_paragraph("Click the Export button on the Table View to download your entire collection as an Excel (.xlsx) or CSV file anytime. Your data is always 100% yours.")

    # Section 10
    add_h1("10. Changing Language")
    doc.add_paragraph("Click your user menu in the top right to switch between 🇬🇧 English and 🇺🇿 O'zbekcha.")

    out_file = DOCS_DIR / "USER_GUIDE.docx"
    doc.save(out_file)
    print(f"Generated DOCX: {out_file}")


def create_pdf():
    pdf_path = DOCS_DIR / "USER_GUIDE.pdf"
    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=12,
        alignment=1, # Center
    )
    h1_style = ParagraphStyle(
        'DocH1',
        parent=styles['Heading2'],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#0e7490'),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True,
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading3'],
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#334155'),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=5,
    )
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=body_style,
        leftIndent=15,
        spaceAfter=3,
    )

    story = []
    story.append(Paragraph("📖 Book Library — User Guide & Quick Start", title_style))
    story.append(Paragraph("Welcome to your personal and household <b>Book Library</b>! Whether you have 20 books or 2,000, this app helps you catalog your physical collection, track what you are reading, lend books to friends without losing them, and decide what to read next.", body_style))
    story.append(Spacer(1, 8))

    def add_rl_img(filename, width=420, height=200):
        img_p = IMAGES_DIR / filename
        if img_p.exists():
            story.append(RLImage(str(img_p), width=width, height=height))
            story.append(Spacer(1, 6))

    # 1. Viewing
    story.append(Paragraph("1. Viewing & Exploring Your Library", h1_style))
    add_rl_img("gallery_view.png", width=480, height=190)
    story.append(Paragraph("<b>🖼️ Gallery View:</b> Displays books with cover art. Click the status badge on any card to update progress, or click the stars to rate.", body_style))
    story.append(Paragraph("<b>📋 Table View:</b> Spreadsheet view for sorting by author, year, page count, or shelf.", body_style))
    story.append(Paragraph("<b>🏷️ Shelves & Tags:</b> Shelves represent physical rooms (Living Room, Desk). Tags represent topics (Favorites, Sci-Fi).", body_style))

    # 2. Adding Books
    story.append(Paragraph("2. Adding Books (4 Fast Ways)", h1_style))
    story.append(Paragraph("<b>📸 Method A: Shelf Photo Scanner</b> — Snap a photo of your bookshelf; AI reads book spines and metadata automatically.", body_style))
    add_rl_img("add_shelf_photo.png", width=400, height=180)
    
    story.append(Paragraph("<b>📷 Method B: Barcode / ISBN Lookup</b> — Scan the back cover barcode with your camera or type the ISBN to auto-fill book details.", body_style))
    add_rl_img("add_barcode_isbn.png", width=400, height=180)

    story.append(Paragraph("<b>✍️ Method C: Manual Entry</b> — Fill in custom details and upload your own cover photo.", body_style))
    add_rl_img("add_manual_entry.png", width=400, height=180)

    story.append(Paragraph("<b>📁 Method D: CSV / Goodreads Import</b> — Import your reading history from Goodreads in one upload.", body_style))
    add_rl_img("add_csv_goodreads.png", width=400, height=180)

    # 3-10
    story.append(Paragraph("3. Tracking Your Reading Progress & Ratings", h1_style))
    story.append(Paragraph("Mark books as <b>⏳ Unread</b>, <b>📖 Reading</b>, <b>✅ Finished</b>, or <b>🛑 Abandoned</b> directly on book cards with 1 tap.", body_style))

    story.append(Paragraph("4. Sharing with Family & Roommates (Household Library)", h1_style))
    story.append(Paragraph("Invite family members using your 6-character Invite Code. Everyone shares the same physical library but keeps independent reading statuses and private ratings.", body_style))

    story.append(Paragraph("5. Lending Tracker — Never Lose a Book", h1_style))
    story.append(Paragraph("Track who borrowed your book, due dates, and return status under the <b>Book Loans</b> tab.", body_style))

    story.append(Paragraph("6. Reading Goals & Personal Stats", h1_style))
    story.append(Paragraph("Set yearly goals and monitor your live reading pace (Ahead of schedule / On track / Behind pace) on the <b>Stats & Goals</b> tab.", body_style))

    story.append(Paragraph("7. AI 'What to Read Next?' Assistant", h1_style))
    story.append(Paragraph("Get instant book recommendations from your unread shelf tailored to your mood.", body_style))

    story.append(Paragraph("8. Public Shareable Profile", h1_style))
    story.append(Paragraph("Enable your public shelf link (e.g. <code>/share/your-name</code>) to share your collection with friends.", body_style))

    story.append(Paragraph("9. Exporting Your Data", h1_style))
    story.append(Paragraph("Download your entire collection as an Excel (<code>.xlsx</code>) or CSV file anytime from the Table View.", body_style))

    story.append(Paragraph("10. Changing Language", h1_style))
    story.append(Paragraph("Switch seamlessly between 🇬🇧 <b>English</b> and 🇺🇿 <b>O'zbekcha</b> from the user menu.", body_style))

    doc.build(story)
    print(f"Generated PDF: {pdf_path}")


if __name__ == "__main__":
    create_docx()
    create_pdf()
